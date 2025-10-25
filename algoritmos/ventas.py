import openpyxl
from pathlib import Path
from tkinter import messagebox

archivo_excel = Path(__file__).parent / 'archivos' / 'datos.xlsx'

#Listar Ventas
def listar_ventas():
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["ventas"]
    ventas = []

    for numero_fila in range(2, hoja.max_row + 1 ):
        venta = {
            "codigo producto": hoja.cell(row=numero_fila, column=1).value,
            "codigo cliente": hoja.cell(row=numero_fila, column=2).value,
            "cantidad producto": hoja.cell(row=numero_fila, column=3).value,
            "total ventas": hoja.cell(row=numero_fila, column=4).value,
        }
        ventas.append(venta)
    libro.close()
    return ventas

def crear_venta(codigo_producto, codigo_cliente, cantidad_producto, total_ventas):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja_ventas = libro["ventas"]
    hoja_inventario = libro["inventario"]

    codigo_producto = str(codigo_producto).strip()
    encontrado = False

    # Buscar producto en inventario
    for fila in hoja_inventario.iter_rows(min_row=2):
        codigo = str(fila[0].value).strip()
        existencia = fila[2].value

        if codigo == codigo_producto:
            encontrado = True
            if existencia < cantidad_producto:
                libro.close()
                messagebox.showerror("Sin existencias", f"No hay suficientes existencias del producto {codigo_producto}")
                return False
            # Descontar del inventario
            fila[2].value = existencia - cantidad_producto
            break  # Detener búsqueda, ya encontramos el producto

    if not encontrado:
        libro.close()
        messagebox.showerror("Error", f"El producto con {codigo_producto} no existe en el inventario")
        return False

    # Registrar la venta (fuera del bucle)
    proxima_fila = hoja_ventas.max_row + 1
    hoja_ventas.cell(row=proxima_fila, column=1).value = codigo_producto
    hoja_ventas.cell(row=proxima_fila, column=2).value = codigo_cliente
    hoja_ventas.cell(row=proxima_fila, column=3).value = cantidad_producto
    hoja_ventas.cell(row=proxima_fila, column=4).value = total_ventas

    libro.save(archivo_excel)
    libro.close()

#Anular venta
def anular_venta(codigo_producto):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["ventas"]

    for i, fila in enumerate(hoja.iter_rows(min_row=2), start=2):
        if str(fila[0].value).strip() == str(codigo_producto).strip():
            hoja.delete_rows(i)
            libro.save(archivo_excel)
            libro.close()
            return True #Venta Eliminada
    libro.close()
    return False #No fue encontrado