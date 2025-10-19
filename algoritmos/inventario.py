import openpyxl
from pathlib import Path

archivo_excel = Path(__file__).parent / 'archivos' / 'datos.xlsx'

#Listar todos los productos
def listar_productos():
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["inventario"]
    productos = []

    for numero_fila in range(2, hoja.max_row + 1):
        producto = {
            "codigo": hoja.cell(row=numero_fila, column=1).value,
            "nombre": hoja.cell(row=numero_fila, column=2).value,
            "existencia": hoja.cell(row=numero_fila, column=3).value,
            "proveedor": hoja.cell(row=numero_fila, column=4).value,
            "precio": hoja.cell(row=numero_fila, column=5).value,
        }
        productos.append(producto)
    libro.close()
    return productos

#Crear un nuevo producto 
def crear_producto(codigo, nombre, existencia, proveedor, precio):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["inventario"]

    proxima_fila = hoja.max_row + 1

    hoja.cell(row=proxima_fila, column=1).value = codigo
    hoja.cell(row=proxima_fila, column=2).value = nombre
    hoja.cell(row=proxima_fila, column=3).value = existencia
    hoja.cell(row=proxima_fila, column=4).value = proveedor
    hoja.cell(row=proxima_fila, column=5).value = precio

    libro.save(archivo_excel)
    libro.close()

#Actualizar productos 
def actualizar_producto(codigo_producto, nuevo_nombre=None, nueva_existencia=None, nuevo_proveedor=None, nuevo_precio=None):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["inventario"]

    for fila in hoja.iter_rows(min_row=2):
        valor_codigo = str(fila[0].value).strip()
        if valor_codigo == str(codigo_producto).strip():
            if nuevo_nombre is not None:
                fila[1].value = nuevo_nombre
            if nueva_existencia is not None:
                fila[2].value = nueva_existencia
            if nuevo_proveedor is not None:
                fila[3].value = nuevo_proveedor
            if nuevo_precio is not None:
                fila[4].value = nuevo_precio

            libro.save(archivo_excel)
            libro.close()
            return True
    libro.close()
    return False

#Editar existencia del producto
def editar_existencia_producto(codigo_producto, cantidad):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["inventario"]

    for fila in hoja.iter_rows(min_row=2):
        if str(fila[0].value).strip() == str(codigo_producto).strip():
            nueva_existencia = fila[2].value + cantidad
            if nueva_existencia < 0:
                libro.close()
                return False  # No se puede tener existencia negativa
            fila[2].value = nueva_existencia
            libro.save(archivo_excel)
            libro.close()
            return True
    libro.close()
    return False  # Producto no encontrado

#Eliminar producto
def eliminar_producto(codigo_producto):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["inventario"]

    for i, fila in enumerate(hoja.iter_rows(min_row=2), start=2):
        if str(fila[0].value).strip() == str(codigo_producto).strip():
            hoja.delete_rows(i)
            libro.save(archivo_excel)
            libro.close()
            return True
    libro.close()
    return False