import openpyxl
from pathlib import Path

archivo_excel = Path(__file__).parent / 'archivos' / 'datos.xlsx'

#Listar Clientes
def listar_clientes():
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["clientes"]
    clientes = []

    for numero_fila in range(2, hoja.max_row + 1):
        cliente = {
            "codigo": hoja.cell(row=numero_fila, column=1).value, 
            "nombre": hoja.cell(row=numero_fila, column=2).value, 
            "direccion": hoja.cell(row=numero_fila, column=3).value, 
        }
        clientes.append(cliente)
    return clientes

#Crear un cliente nuevo
def crear_cliente(codigo, nombre, dirreccion):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["clientes"]

    proxima_fila = hoja.max_row + 1

    hoja.cell(row=proxima_fila, column=1).value = codigo
    hoja.cell(row=proxima_fila, column=2).value = nombre
    hoja.cell(row=proxima_fila, column=3).value = dirreccion

    libro.save(archivo_excel)
    libro.close()

#Editar cliente
def actualizar_cliente(codigo_cliente, nuevo_nombre, nueva_direccion):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["clientes"]

    for fila in hoja.iter_rows(min_row=2):

        valor_codigo = str(fila[0].value).strip() 
        if valor_codigo == str(codigo_cliente).strip():
            cliente_anterior = {
                "nombre": fila[1].value,
                "direccion": fila[2].value
            }

            if nuevo_nombre:
                fila[1].value = nuevo_nombre
            if nueva_direccion:
                fila[2].value = nueva_direccion

            libro.save(archivo_excel)
            libro.close()
            return True, cliente_anterior
    libro.close()
    return False

#Eliminar cliente
def eliminar_cliente(codigo_cliente):
    libro = openpyxl.load_workbook(archivo_excel)
    hoja = libro["clientes"]

    for fila in hoja.iter_rows(min_row=2):

        valor_codigo = str(fila[0].value).strip() 
        if valor_codigo == str(codigo_cliente).strip():
            hoja.delete_rows(fila[0].row, 1)
            libro.save(archivo_excel)
            libro.close()
            return True #Cliente eliminado
    libro.close()
    return False #Cliente no encontrado