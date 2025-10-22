from tkinter import messagebox
from openpyxl import load_workbook, Workbook
from pathlib import Path


archivo_excel = Path(__file__).parent / 'archivos' / 'datos.xlsx'

def cargar_datos():
    libro = load_workbook(archivo_excel, data_only=True)
    clientes_hoja = libro['clientes']
    ventas_hoja = libro['ventas']
    inventario_hoja = libro['inventario']

    # Cargar clientes en diccionario {codigo: {nombre, direccion}}
    clientes = {}
    for fila in clientes_hoja.iter_rows(min_row=2, max_col=3, values_only=True):
        codigo, nombre, direccion = fila
        clientes[codigo] = {
            "nombre": nombre,
            "direccion": direccion
        }

    # Cargar productos
    productos = {}
    for fila in inventario_hoja.iter_rows(min_row=2, max_col=2, values_only=True):
        codigo_producto, nombre_producto = fila
        productos[codigo_producto] = nombre_producto

    # Cargar ventas en lista de diccionarios
    ventas = []
    for fila in ventas_hoja.iter_rows(min_row=2, max_col=4, values_only=True):
        codigo_producto, codigo_cliente, cantidad_productos, total_venta = fila
        ventas.append({
            "codigo_producto": codigo_producto,
            "codigo_cliente": codigo_cliente,
            "cantidad_productos": cantidad_productos,
            "total_venta": total_venta
        })

    return clientes, productos, ventas

def ventas_por_cliente():
    clientes, _, ventas = cargar_datos()
    resumen = {}
    for v in ventas:
        codigo_cliente = v["codigo_cliente"]
        resumen[codigo_cliente] = resumen.get(codigo_cliente, 0) + (v["total_venta"] or 0)

    # Retorna lista de tuplas: (codigo_cliente, nombre_cliente, total_venta)
    return [(codigo, clientes.get(codigo, {}).get("nombre", "Desconocido"), total) for codigo, total in resumen.items()]

# def ventas_por_producto():
#     _, ventas = cargar_datos()
#     resumen = {}

#     for v in ventas:
#         codigo_producto = v["codigo_producto"]
#         cantidad = v["cantidad_productos"] or 0
#         total = v["total_venta"] or 0

#         if codigo_producto not in resumen:
#             resumen[codigo_producto] = {"cantidad": 0, "total": 0}

#         resumen[codigo_producto]["cantidad"] += cantidad
#         resumen[codigo_producto]["total"] += total

#     # Crear lista de tuplas (codigo_producto, cantidad_total, total_ventas)
#     resultado = [
#         (codigo, datos["cantidad"], datos["total"]) for codigo, datos in resumen.items()
#     ]

#     # Opcional: ordenar por código de producto
#     return sorted(resultado)


def ventas_por_producto():
    _, productos, ventas = cargar_datos()
    resumen = {}

    for v in ventas:
        codigo_producto = v["codigo_producto"]
        cantidad = v["cantidad_productos"] or 0
        total = v["total_venta"] or 0

        if codigo_producto not in resumen:
            resumen[codigo_producto] = {"cantidad": 0, "total": 0}

        resumen[codigo_producto]["cantidad"] += cantidad
        resumen[codigo_producto]["total"] += total

    resultado = []
    for codigo, datos in resumen.items():
        nombre = productos.get(codigo, "Desconocido")
        resultado.append((codigo, nombre, datos["cantidad"], datos["total"]))

    return sorted(resultado)


#Guardamos el reporte en libro excel 
def guardar_reporte_ventas_por_cliente(nombre_archivo="ventas_por_cliente.xlsx"):
    datos = ventas_por_cliente() 

    # Crear libro y hoja
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas por Cliente"

    hoja.append(["Código Cliente", "Nombre Cliente", "Total Venta"])

    for codigo, nombre, total in datos:
        hoja.append([codigo, nombre, total])

    libro.save(nombre_archivo)
    messagebox.showinfo("Éxito", f"Reporte guardado con éxito")

#Ventas por producto

def guardar_reporte_ventas_por_producto(nombre_archivo="ventas_por_producto.xlsx"):
    datos = ventas_por_producto() 

    # Crear libro y hoja
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas por Producto"

    hoja.append(["Código producto", "Nombre Producto", "Cantidad Total", "Total Venta"])

    for codigo, nombre, cantidad, total in datos:
        hoja.append([codigo, nombre, cantidad, total])

    libro.save(nombre_archivo)
    messagebox.showinfo("Éxito", f"Reporte guardado con éxito")

